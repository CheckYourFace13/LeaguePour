"""Builds the free VenueSprocket template downloads in public/downloads/.

    python scripts/generate-template-downloads.py   (needs: pip install python-docx reportlab)

Content mirrors the pages it belongs to - keep them in sync when either changes:
  BEO      -> src/app/venuesprocket/guides/beo-template/page.tsx
  Proposal -> src/app/venuesprocket/templates/event-proposal-template/page.tsx
Every name, number, and email in the examples is fictional (555 numbers, example.com).
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor, Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Flowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent.parent / "public" / "downloads"
FOOTER = "Free template from VenueSprocket - venuesprocket.com"
SAMPLE_NOTE = "Sample only: the venue, host, and contact details below are fictional."

# ---------------------------------------------------------------------------------------- BEO
BEO_TITLE = "Banquet Event Order (BEO)"
BEO_INTRO = ("Fill in one BEO per event and give a copy to every team working it - kitchen, bar, floor, "
             "and management. Fields marked (required) should be filled in for every event.")
BEO_SECTIONS = [
    ("Event basics", [
        ("Event name", True), ("Event date", True), ("Venue open / setup start time", True),
        ("Guest arrival time", True), ("Event end / teardown complete time", True),
        ("Room or space assigned", True), ("Guest count (confirmed)", True),
        ("Guest count (guaranteed minimum, if applicable)", False),
    ]),
    ("Contact information", [
        ("Client / host name", True), ("Client phone number", True), ("Client email", True),
        ("On-site day-of contact (if different from client)", False),
        ("Internal event manager / coordinator", True),
    ]),
    ("Food & beverage", [
        ("Menu / package selected", True),
        ("Beverage package or bar setup (open bar, cash bar, drink tickets, beer/wine only)", True),
        ("Food and beverage minimum (if applicable)", False),
        ("Allergies / dietary restrictions (write 'none reported' rather than leaving blank)", True),
        ("Cake / dessert service details", False),
        ("Course timing (when each course should be served)", False),
    ]),
    ("Room setup & AV", [
        ("Table layout (rounds, banquet-style, cocktail rounds, etc.)", True),
        ("Seating chart or assigned seating notes", False),
        ("AV equipment needed (mic, speaker, projector, screen)", False),
        ("Decor restrictions (candles, confetti, adhesives on walls)", False),
        ("Vendor access (DJ, photographer, florist) - arrival time and load-in notes", False),
    ]),
    ("Staffing", [
        ("Servers / bartenders assigned to this event", True), ("Manager on duty", True),
        ("Kitchen lead for this event", False),
    ]),
    ("Payment status & internal notes", [
        ("Deposit paid (amount, date)", True), ("Balance due and due date", True),
        ("Internal notes - staff only, never shown to the customer", False),
    ]),
]
BEO_EXAMPLE = [
    ("Event basics", [
        ("Event", "Marino 40th Birthday Buyout"), ("Date", "Saturday, October 17"),
        ("Times", "Setup 5:00 PM  |  Guest arrival 7:00 PM  |  Event end 11:00 PM"),
        ("Room", "Taproom back patio (full buyout)"), ("Guest count", "45 confirmed"),
    ]),
    ("Contact", [
        ("Client", "Dana Marino  |  (555) 019-2244  |  dana.marino@example.com"),
        ("Day-of contact", "Same as client"), ("Event manager", "James (assistant manager)"),
    ]),
    ("Food & beverage", [
        ("Menu", "Taco bar package + 2 shareable appetizer trays"),
        ("Bar", "Open bar, beer/wine/well only, 4-hour limit (cuts off 11:00 PM)"),
        ("F&B minimum", "$1,800"), ("Allergies", "One guest - tree nut allergy, kitchen notified"),
        ("Cake", "Client bringing own cake; staff hold it in the walk-in until the 8:30 PM cut"),
    ]),
    ("Room setup & AV", [
        ("Setup", "Patio in lounge configuration, 6 cocktail rounds + bar seating"),
        ("AV", "Client bringing a bluetooth speaker; staff to confirm patio speaker pairing"),
        ("Decor", "No confetti (patio policy) - balloons and banner OK"), ("Vendors", "None"),
    ]),
    ("Staffing", [
        ("Staff", "Priya + Marcus (bar), Deshawn (floor)"), ("Manager on duty", "James"),
        ("Kitchen lead", "Ana"),
    ]),
    ("Payment & notes", [
        ("Deposit", "$400 paid 9/12"), ("Balance due", "$1,400 (F&B minimum less deposit), due night of event"),
        ("Internal notes", "Repeat customer (hosted an anniversary party in June)."),
    ]),
]
BEO_CHECKLIST = [
    "Guest count reconfirmed with the client within the last week",
    "Menu and bar package match the signed contract - no last-minute changes missed",
    "Allergies and dietary restrictions confirmed with the kitchen directly, not just noted on paper",
    "Room setup matches what the client walked through or approved",
    "AV/equipment tested if the event depends on it (mic, speaker, screen)",
    "Staffing assigned and confirmed - and told what this specific event needs",
    "Deposit and balance-due amounts match your records exactly",
    "BEO distributed to every team involved (kitchen, bar, floor, management)",
]

# ----------------------------------------------------------------------------------- Proposal
PROP_TITLE = "Private Event Proposal"
PROP_INTRO = ("Replace the bracketed text and send. Keep it to one page, show how the price is built, "
              "and end with a single next step.")
PROP_SECTIONS = [
    ("Prepared for / by", [
        ("Prepared for", "[Host name]  |  [Phone]  |  [Email]"),
        ("Prepared by", "[Venue name]  |  [Your name]  |  [Date sent]"),
    ]),
    ("The event", [
        ("Event", "[Occasion, e.g. \"Rivera 40th birthday\"]"), ("Date", "[Day, date]"),
        ("Time", "[Arrival time] to [End time]"), ("Space", "[Room or area]  |  [Private room / full buyout]"),
        ("Guests", "[Expected count]  |  Final count due [date]"),
    ]),
    ("Pricing", [
        ("Room fee", "$[amount]"), ("Food & beverage minimum", "$[amount]  (before tax and service charge)"),
        ("Service charge / gratuity", "[__]%"), ("Sales tax", "[__]%"),
        ("Estimated total", "$[amount]  (estimate - final bill depends on what is ordered)"),
    ]),
    ("To reserve the date", [
        ("Deposit due at signing", "$[amount]  (applied to your final bill)"),
        ("Balance due", "[At the end of the event / __ days before]"),
    ]),
    ("Key terms", [
        ("Cancellation", "[One sentence, e.g. \"Deposit refundable until 30 days before the event.\"]"),
        ("Rescheduling", "[One sentence, e.g. \"One date change within 6 months, subject to availability.\"]"),
    ]),
]
PROP_CLOSE = ["Full terms are in the event contract.",
              "This proposal is valid until [date]. The date is not held until the deposit is paid.",
              "Next step: reply to accept, sign the contract, and pay the deposit."]
PROP_EXAMPLE = [
    ("Prepared for / by", [
        ("Prepared for", "Sam Rivera  |  (555) 010-4477  |  sam.rivera@example.com"),
        ("Prepared by", "The Juniper Room  |  Alex, events  |  March 3"),
    ]),
    ("The event", [
        ("Event", "Rivera 40th birthday"), ("Date", "Saturday, April 18"), ("Time", "6:30 PM to 10:30 PM"),
        ("Space", "Back dining room  |  Private room"), ("Guests", "35 expected  |  Final count due April 11"),
    ]),
    ("Pricing", [
        ("Room fee", "$200"), ("Food & beverage minimum", "$1,750  (before tax and service charge)"),
        ("Service charge / gratuity", "20%"), ("Sales tax", "8%"),
        ("Estimated total", "$2,440  (estimate - $200 + $1,750 + $350 service + $140 tax)"),
    ]),
    ("To reserve the date", [
        ("Deposit due at signing", "$500  (applied to your final bill)"), ("Balance due", "At the end of the event"),
    ]),
    ("Key terms", [
        ("Cancellation", "Deposit refundable until 30 days before the event."),
        ("Rescheduling", "One date change within 6 months, subject to availability."),
    ]),
]
PROP_EXAMPLE_CLOSE = ["Full terms are in the event contract.",
                      "This proposal is valid until March 10. The date is not held until the deposit is paid.",
                      "Next step: reply to accept, sign the contract, and pay the deposit."]
PROP_CHECKLIST = [
    "Host name, phone, and email",
    "Date, arrival time, end time, and the room or area being booked",
    "Expected guest count and the final-count deadline",
    "Minimum labeled as before or after tax and service charge",
    "Estimated total labeled as an estimate",
    "Deposit amount, balance due date, and proposal expiry date",
    "One clear next step: accept, sign the contract, pay the deposit",
]

ACCENT = RGBColor(0x1A, 0x5F, 0x3F)


# ------------------------------------------------------------------------------------- DOCX
def docx_base(title):
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.footer.paragraphs[0].text = FOOTER
        s.footer.paragraphs[0].runs[0].font.size = Pt(8)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    h = d.add_heading(title, level=0)
    h.runs[0].font.color.rgb = ACCENT
    return d


def docx_table(d, rows, label_w=2.4, value_w=4.5, blank=False):
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Inches(label_w), Inches(value_w)
        r = cells[0].paragraphs[0].add_run(label)
        r.bold = True
        r.font.size = Pt(9.5)
        cells[1].paragraphs[0].add_run("" if blank else value).font.size = Pt(10)
    d.add_paragraph()


def docx_checklist(d, title, items):
    d.add_heading(title, level=2)
    for item in items:
        d.add_paragraph("☐  " + item)


def build_beo_docx():
    d = docx_base(BEO_TITLE)
    d.add_paragraph(BEO_INTRO)
    for title, fields in BEO_SECTIONS:
        d.add_heading(title, level=2)
        docx_table(d, [(f"{label}{' (required)' if req else ''}", "") for label, req in fields], 3.0, 3.9, blank=True)
    docx_checklist(d, "Before you hand it out", BEO_CHECKLIST)
    d.add_page_break()
    d.add_heading("Example: a filled-in BEO", level=1)
    d.add_paragraph(SAMPLE_NOTE).runs[0].italic = True
    for title, rows in BEO_EXAMPLE:
        d.add_heading(title, level=2)
        docx_table(d, rows, 1.6, 5.3)
    d.save(OUT / "venuesprocket-beo-template.docx")


def build_proposal_docx():
    d = docx_base(PROP_TITLE)
    d.add_paragraph(PROP_INTRO)
    for title, rows in PROP_SECTIONS:
        d.add_heading(title, level=2)
        docx_table(d, rows)
    for line in PROP_CLOSE:
        d.add_paragraph(line)
    d.add_page_break()
    d.add_heading("Example: a filled-in proposal", level=1)
    d.add_paragraph(SAMPLE_NOTE).runs[0].italic = True
    for title, rows in PROP_EXAMPLE:
        d.add_heading(title, level=2)
        docx_table(d, rows)
    for line in PROP_EXAMPLE_CLOSE:
        d.add_paragraph(line)
    docx_checklist(d, "Before you send (not part of the proposal)", PROP_CHECKLIST)
    d.save(OUT / "venuesprocket-event-proposal-template.docx")


# -------------------------------------------------------------------------------------- PDF
ss = getSampleStyleSheet()
H0 = ParagraphStyle("h0", parent=ss["Title"], alignment=0, textColor=colors.HexColor("#1a5f3f"), fontSize=22, spaceAfter=8)
H2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#333333"))
H2S = ParagraphStyle("h2s", parent=H2, spaceBefore=6, spaceAfter=3)
BODY = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13)
CELL = ParagraphStyle("c", parent=BODY, fontSize=9, leading=12)
CELLB = ParagraphStyle("cb", parent=CELL, fontName="Helvetica-Bold")
NOTE = ParagraphStyle("n", parent=BODY, fontName="Helvetica-Oblique", textColor=colors.HexColor("#666666"))


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_table(rows, label_w=2.3, blank=False, row_h=None, pad=4):
    data = [[Paragraph(esc(l), CELLB), Paragraph("" if blank else esc(v), CELL)] for l, v in rows]
    t = Table(data, colWidths=[label_w * inch, (6.9 - label_w) * inch], rowHeights=row_h)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f2f5f3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), pad), ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
    ]))
    return t


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#888888"))
    canvas.drawString(0.8 * inch, 0.45 * inch, FOOTER)
    canvas.drawRightString(7.7 * inch, 0.45 * inch, f"Page {doc.page}")
    canvas.restoreState()


def pdf_doc(path, title):
    return SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                             topMargin=0.7 * inch, bottomMargin=0.8 * inch, title=title, author="VenueSprocket")


class TickBox(Flowable):
    """A small empty square to tick by hand (Helvetica has no ballot-box glyph)."""
    def __init__(self, size=8.5):
        super().__init__()
        self.width = self.height = size

    def draw(self):
        self.canv.setStrokeColor(colors.HexColor("#555555"))
        self.canv.setLineWidth(0.8)
        self.canv.rect(0, -1.5, self.width, self.height)


def checklist_flow(title, items):
    t = Table([[TickBox(), Paragraph(esc(i), BODY)] for i in items], colWidths=[0.2 * inch, 6.6 * inch])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 2),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    return [KeepTogether([Paragraph(title, H2), t])]


def build_beo_pdf():
    f = [Paragraph(BEO_TITLE, H0), Paragraph(esc(BEO_INTRO), BODY)]
    for title, fields in BEO_SECTIONS:
        f += [Paragraph(esc(title), H2),
              pdf_table([(f"{l}{' (required)' if r else ''}", "") for l, r in fields], 3.0, blank=True, row_h=None)]
    f += checklist_flow("Before you hand it out", BEO_CHECKLIST)
    f += [PageBreak(), Paragraph("Example: a filled-in BEO", H0), Paragraph(SAMPLE_NOTE, NOTE)]
    for title, rows in BEO_EXAMPLE:
        f += [Paragraph(esc(title), H2S), pdf_table(rows, 1.6, pad=3)]
    pdf_doc(OUT / "venuesprocket-beo-template.pdf", "BEO template").build(f, onFirstPage=footer, onLaterPages=footer)


def build_proposal_pdf():
    f = [Paragraph(PROP_TITLE, H0), Paragraph(esc(PROP_INTRO), BODY)]
    for title, rows in PROP_SECTIONS:
        f += [Paragraph(esc(title), H2), pdf_table(rows)]
    f += [Spacer(1, 6)] + [Paragraph(esc(l), BODY) for l in PROP_CLOSE]
    f += [PageBreak(), Paragraph("Example: a filled-in proposal", H0), Paragraph(SAMPLE_NOTE, NOTE)]
    for title, rows in PROP_EXAMPLE:
        f += [Paragraph(esc(title), H2S), pdf_table(rows, pad=3)]
    f += [Spacer(1, 6)] + [Paragraph(esc(l), BODY) for l in PROP_EXAMPLE_CLOSE]
    f += checklist_flow("Before you send (not part of the proposal)", PROP_CHECKLIST)
    pdf_doc(OUT / "venuesprocket-event-proposal-template.pdf", "Event proposal template").build(f, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_beo_docx(); build_beo_pdf(); build_proposal_docx(); build_proposal_pdf()
    for p in sorted(OUT.iterdir()):
        print(p.name, p.stat().st_size)
